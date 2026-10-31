"""
Ustozlar va Repetitorlar Kabineti — Guruh Tahlili, Monitoring va Excel/PDF Eksport Xizmati.
"""

from __future__ import annotations

import io
import logging
from datetime import datetime
from typing import Any

from django.db.models import Avg, Count, Max, Q
from django.utils import timezone

from apps.accounts.models import User
from apps.courses.models import StudentGroup
from apps.essays.models import EssaySubmission
from apps.results.models import Result
from apps.tests.bmb_grading import calculate_bmb_grade

logger = logging.getLogger(__name__)


def get_group_full_analytics(group: StudentGroup) -> dict[str, Any]:
    """
    Guruh bo'yicha to'liq jonli monitoring:
      - Har bir o'quvchining test va esse statistikasi
      - Rasmiy BMB darajalari (A+, A, B+, B, C+, C, Fail)
      - Guruh reytingi va taqsimoti
    """
    students = list(
        group.students.filter(is_active=True).order_by("first_name", "last_name")
    )
    student_ids = [s.id for s in students]

    # Test natijalari
    test_results = (
        Result.objects.filter(student_id__in=student_ids)
        .select_related("test")
        .order_by("-calculated_at")
    )
    results_by_student: dict[int, list[Result]] = {}
    for r in test_results:
        results_by_student.setdefault(r.student_id, []).append(r)

    # Esse natijalari
    essay_subs = (
        EssaySubmission.objects.filter(student_id__in=student_ids)
        .exclude(status=EssaySubmission.Status.DRAFT)
        .order_by("-updated_at")
    )
    essays_by_student: dict[int, list[EssaySubmission]] = {}
    for es in essay_subs:
        essays_by_student.setdefault(es.student_id, []).append(es)

    # Kurs to'lovlari (Tuition Billing)
    from apps.tuition.models import StudentTuition

    tuitions = StudentTuition.objects.filter(
        group=group, student_id__in=student_ids
    ).order_by("-due_date")
    tuition_by_student: dict[int, StudentTuition] = {}
    for tuit in tuitions:
        if tuit.student_id not in tuition_by_student:
            tuition_by_student[tuit.student_id] = tuit

    roster = []
    bmb_counts = {"A+": 0, "A": 0, "B+": 0, "B": 0, "C+": 0, "C": 0, "Fail": 0}
    all_pcts: list[float] = []

    for s in students:
        s_results = results_by_student.get(s.id, [])
        s_essays = essays_by_student.get(s.id, [])

        tests_count = len(s_results)
        avg_test_pct = (
            round(sum(float(r.percentage) for r in s_results) / tests_count, 1)
            if tests_count > 0
            else 0.0
        )
        if tests_count > 0:
            all_pcts.append(avg_test_pct)

        latest_test = s_results[0] if s_results else None
        latest_score_str = (
            f"{latest_test.score}/{latest_test.max_score}" if latest_test else "—"
        )
        latest_pct = float(latest_test.percentage) if latest_test else 0.0

        # BMB darajasi
        bmb_info = calculate_bmb_grade(avg_test_pct if tests_count > 0 else 0.0)
        grade = bmb_info["grade"]
        if tests_count > 0 and grade in bmb_counts:
            bmb_counts[grade] += 1

        # Esse ballari
        essays_count = len(s_essays)
        graded_essays = [e for e in s_essays if e.total_score is not None]
        avg_essay_score = (
            round(
                sum(
                    float(e.final_score if e.final_score is not None else e.total_score)
                    for e in graded_essays
                )
                / len(graded_essays),
                1,
            )
            if graded_essays
            else 0.0
        )

        last_active = None
        dates = []
        if latest_test:
            dates.append(latest_test.calculated_at)
        if s_essays:
            dates.append(s_essays[0].updated_at)
        if dates:
            last_active = max(dates).strftime("%d.%m.%Y %H:%M")

        roster.append(
            {
                "student_id": s.id,
                "student_name": s.get_full_name() or s.email,
                "full_name": s.get_full_name() or s.email,
                "email": s.email,
                "has_pro": True,
                "pro_status_label": "⭐ PRO (Faol)",
                "tests_count": tests_count,
                "tests_taken": tests_count,
                "avg_test_percentage": avg_test_pct,
                "avg_test_pct": avg_test_pct,
                "latest_test_score": latest_score_str,
                "latest_test_pct": latest_pct,
                "latest_test_grade": grade if tests_count > 0 else "—",
                "essays_count": essays_count,
                "essays_written": essays_count,
                "avg_essay_score": avg_essay_score,
                "bmb_grade": grade if tests_count > 0 else "—",
                "bmb_title": bmb_info["title"]
                if tests_count > 0
                else "Test topshirmagan",
                "badge_class": bmb_info["badge_class"] if tests_count > 0 else "badge",
                "tuition_status": (
                    tuition_by_student.get(s.id).status
                    if tuition_by_student.get(s.id)
                    else "pending"
                ),
                "tuition_label": (
                    "To'langan"
                    if tuition_by_student.get(s.id)
                    and tuition_by_student.get(s.id).status == "paid"
                    else (
                        "Muddati o'tgan"
                        if tuition_by_student.get(s.id)
                        and tuition_by_student.get(s.id).status == "overdue"
                        else "Kutilmoqda"
                    )
                ),
                "last_active": last_active or "Faollik yo'q",
            }
        )

    # Sort roster by avg_test_pct descending
    roster.sort(key=lambda x: x["avg_test_pct"], reverse=True)
    for idx, item in enumerate(roster, start=1):
        item["rank"] = idx

    group_avg_pct = round(sum(all_pcts) / len(all_pcts), 1) if all_pcts else 0.0
    group_bmb = calculate_bmb_grade(group_avg_pct)["grade"] if all_pcts else "—"

    return {
        "group_id": group.id,
        "group_name": group.name,
        "invite_code": group.invite_code or f"GR-{group.id}",
        "teacher_name": group.teacher.get_full_name() or group.teacher.email,
        "total_students": len(students),
        "group_avg_percentage": group_avg_pct,
        "group_bmb_grade": group_bmb,
        "bmb_distribution": bmb_counts,
        "summary": {
            "total_tests_taken": sum(r["tests_taken"] for r in roster),
            "total_essays_submitted": sum(r["essays_written"] for r in roster),
            "average_test_score": group_avg_pct,
            "average_test_percentage": group_avg_pct,
            "average_essay_score": round(
                sum(r["avg_essay_score"] for r in roster) / len(roster), 1
            )
            if roster
            else 0.0,
            "bmb_distribution": bmb_counts,
        },
        "students": roster,
    }


def _activate_student_pro_for_group(student: User, group: StudentGroup) -> None:
    """O'quvchi ustoz guruhiga qo'shilganda unga avtomatik PRO tarifini faollashtiradi."""
    from apps.payments.models import (
        SubscriptionPlan,
        UserSubscription,
        TeacherSponsorship,
    )
    from apps.payments.subscriptions import grant_subscription
    from django.utils import timezone

    # 1. Update Profile plan
    try:
        profile = getattr(student, "profile", None)
        if profile:
            profile.plan = "pro"
            profile.save(update_fields=["plan"])
    except Exception as e:
        logger.warning(
            "Could not update profile plan for student %s: %s", student.id, e
        )

    # 2. Grant 365-day PRO subscription
    pro_plan = SubscriptionPlan.objects.filter(
        plan_type=SubscriptionPlan.PlanType.PRO
    ).first()
    if pro_plan:
        try:
            grant_subscription(
                user=student,
                plan=pro_plan,
                source=UserSubscription.Source.TEACHER_SPONSORED,
                grant_source=f"group:{group.id}",
                days=365,
                reason=f"Ustoz guruhi orqali avtomatik PRO: {group.name} (#{group.invite_code})",
            )
        except Exception as e:
            logger.warning(
                "Could not grant PRO subscription to student %s: %s", student.id, e
            )

    # 3. Create or update approved TeacherSponsorship
    try:
        now = timezone.now()
        TeacherSponsorship.objects.update_or_create(
            student=student,
            group=group,
            defaults={
                "is_active": True,
                "approved_at": now,
                "expires_at": now + timezone.timedelta(days=365),
                "is_unlimited_tests": True,
                "is_unlimited_image_essays": True,
                "is_unlimited_text_essays": True,
                "note": f"Ustoz {group.teacher.get_full_name()} guruhi orqali avtomatik berilgan PRO",
            },
        )
    except Exception as e:
        logger.warning(
            "Could not create teacher sponsorship for student %s: %s", student.id, e
        )


def find_group_by_token_or_code(token_or_code: str) -> StudentGroup | None:
    """Finds a StudentGroup by invite_token or invite_code."""
    if not token_or_code:
        return None
    raw = token_or_code.strip()
    clean_code = raw.removeprefix("#").upper()
    return (
        StudentGroup.objects.filter(invite_token=raw).first()
        or StudentGroup.objects.filter(invite_code__iexact=clean_code).first()
        or StudentGroup.objects.filter(invite_code__iexact=f"#{clean_code}").first()
    )


def join_group_by_token(student: User, token_or_code: str) -> tuple[str, StudentGroup]:
    """
    Joins student to group using invite token or code.
    Returns:
      ("ALREADY_MEMBER", group) if student is already in the group
      ("SUCCESS", group) on successful addition with auto PRO activation.
    Raises ValueError if group is not found.
    """
    group = find_group_by_token_or_code(token_or_code)
    if not group:
        raise ValueError(f"'{token_or_code}' bo'yicha guruh topilmadi.")

    if group.students.filter(id=student.id).exists():
        return ("ALREADY_MEMBER", group)

    group.students.add(student)
    _activate_student_pro_for_group(student, group)
    logger.info("Student %s joined group %s via token %s", student.id, group.id, token_or_code)
    return ("SUCCESS", group)


def join_group_by_invite_code(student: User, invite_code: str) -> StudentGroup:
    """O'quvchini taklif kodi orqali guruhga qo'shadi va avtomatik PRO tarifini yoqadi."""
    raw = (invite_code or "").strip().upper()
    if raw.startswith("#"):
        raw = raw[1:]

    group = (
        StudentGroup.objects.filter(invite_code__iexact=raw).first()
        or StudentGroup.objects.filter(invite_code__iexact=f"#{raw}").first()
    )

    if not group:
        raise ValueError(f"'{invite_code}' kodi bo'yicha guruh topilmadi.")

    if not group.students.filter(id=student.id).exists():
        group.students.add(student)
        logger.info("Student %s joined group %s via code %s", student.id, group.id, raw)

    # Avtomatik PRO tarifini faollashtirish
    _activate_student_pro_for_group(student, group)

    return group


def export_group_to_excel(group: StudentGroup) -> bytes:
    """Guruh monitoringini zamonaviy Excel (.xlsx) fayliga eksport qiladi."""
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    analytics = get_group_full_analytics(group)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Guruh Monitoringi"
    ws.views.sheetView[0].showGridLines = True

    # Ranglar
    HEADER_BG = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    ROW_ALT_BG = PatternFill(
        start_color="F8FAFC", end_color="F8FAFC", fill_type="solid"
    )
    SUMMARY_BG = PatternFill(
        start_color="E2E8F0", end_color="E2E8F0", fill_type="solid"
    )

    FONT_TITLE = Font(name="Calibri", size=16, bold=True, color="0F172A")
    FONT_SUBTITLE = Font(name="Calibri", size=11, italic=True, color="475569")
    FONT_HEADER = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    FONT_REGULAR = Font(name="Calibri", size=11, color="1E293B")
    FONT_BOLD = Font(name="Calibri", size=11, bold=True, color="0F172A")

    THIN_BORDER = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )

    # 1. Sarlavha
    ws.merge_cells("A1:I1")
    ws["A1"] = f"GRAND EDUCATION — {group.name.upper()} GURUHI HISOBOTI"
    ws["A1"].font = FONT_TITLE
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 30

    # 2. Qo'shimcha ma'lumotlar
    ws["A2"] = (
        f"Ustoz: {analytics['teacher_name']}  |  Guruh kodi: {analytics['invite_code']}  |  Sana: {datetime.now().strftime('%d.%m.%Y %H:%M')}"
    )
    ws["A2"].font = FONT_SUBTITLE
    ws.merge_cells("A2:I2")
    ws.row_dimensions[2].height = 20

    # Bo'sh qator
    ws.row_dimensions[3].height = 10

    # 3. Ustunlar sarlavhasi
    headers = [
        "№",
        "O'quvchi F.I.Sh.",
        "Email / Login",
        "Tarif",
        "Testlar soni",
        "O'rtacha ko'rsatkich (%)",
        "BMB Sertifikat Darajasi",
        "Esselar soni",
        "O'rtacha esse bali",
        "To'lov holati",
        "Oxirgi faollik",
    ]

    ws.row_dimensions[4].height = 26
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=4, column=col_idx, value=header)
        cell.fill = HEADER_BG
        cell.font = FONT_HEADER
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
        cell.border = THIN_BORDER

    # 4. Ma'lumotlar qatorlari
    current_row = 5
    for item in analytics["students"]:
        ws.row_dimensions[current_row].height = 22
        is_even = current_row % 2 == 0
        row_fill = ROW_ALT_BG if is_even else None

        row_data = [
            item["rank"],
            item["full_name"],
            item["email"],
            item.get("pro_status_label", "⭐ PRO (Faol)"),
            item["tests_taken"],
            f"{item['avg_test_pct']}%",
            item["bmb_grade"],
            item["essays_written"],
            f"{item['avg_essay_score']} ball",
            item.get("tuition_label", "Kutilmoqda"),
            item["last_active"],
        ]

        for col_idx, val in enumerate(row_data, start=1):
            cell = ws.cell(row=current_row, column=col_idx, value=val)
            cell.font = FONT_REGULAR
            cell.border = THIN_BORDER
            if row_fill:
                cell.fill = row_fill

            # Formatlash
            if col_idx in (1, 4, 5, 7, 8, 10):
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col_idx in (6, 9):
                cell.alignment = Alignment(horizontal="right", vertical="center")
                cell.font = FONT_BOLD
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

        current_row += 1

    # 5. Jami qatori
    ws.row_dimensions[current_row].height = 24
    summary_cells = [
        ("A", "B", "JAMI O'QUVCHILAR:"),
        ("C", "D", f"{analytics['total_students']} nafar"),
        ("E", "F", f"O'rtacha guruh bali: {analytics['group_avg_percentage']}%"),
        ("G", "G", f"BMB: {analytics['group_bmb_grade']}"),
        (
            "H",
            "J",
            f"A+: {analytics['bmb_distribution']['A+']} | A: {analytics['bmb_distribution']['A']} | B+: {analytics['bmb_distribution']['B+']} | B: {analytics['bmb_distribution']['B']}",
        ),
    ]
    for start_c, end_c, val in summary_cells:
        if start_c != end_c:
            ws.merge_cells(f"{start_c}{current_row}:{end_c}{current_row}")
        c = ws[f"{start_c}{current_row}"]
        c.value = val
        c.font = FONT_BOLD
        c.fill = SUMMARY_BG
        c.alignment = Alignment(horizontal="center", vertical="center")

    for col in range(1, 11):
        ws.cell(row=current_row, column=col).border = THIN_BORDER

    # Ustun kengliklarini avtomatik sozlash
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_group_to_pdf(group: StudentGroup) -> bytes:
    """Guruh hisobotini chiroyli PDF formatida yaratadi (ReportLab)."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import (
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    analytics = get_group_full_analytics(group)

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=landscape(A4),
        leftMargin=1.5 * cm,
        rightMargin=1.5 * cm,
        topMargin=1.5 * cm,
        bottomMargin=1.5 * cm,
    )

    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        "DocTitle",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=18,
        leading=22,
        textColor=colors.HexColor("#0F172A"),
    )
    meta_style = ParagraphStyle(
        "DocMeta",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=10,
        leading=14,
        textColor=colors.HexColor("#475569"),
    )
    th_style = ParagraphStyle(
        "TH",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=9,
        leading=11,
        textColor=colors.white,
        alignment=1,  # center
    )
    td_style = ParagraphStyle(
        "TD",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=8.5,
        leading=11,
        textColor=colors.HexColor("#1E293B"),
    )
    td_bold = ParagraphStyle(
        "TDBold",
        parent=td_style,
        fontName="Helvetica-Bold",
        alignment=1,
    )

    story = []

    # Sarlavha
    story.append(
        Paragraph(
            f"GRAND EDUCATION — {group.name.upper()} GURUHI HISOBOTI", title_style
        )
    )
    story.append(Spacer(1, 4))
    meta_text = (
        f"<b>Ustoz:</b> {analytics['teacher_name']} &nbsp;&nbsp;|&nbsp;&nbsp; "
        f"<b>Guruh kodi:</b> {analytics['invite_code']} &nbsp;&nbsp;|&nbsp;&nbsp; "
        f"<b>Jami o'quvchilar:</b> {analytics['total_students']} nafar &nbsp;&nbsp;|&nbsp;&nbsp; "
        f"<b>O'rtacha ball:</b> {analytics['group_avg_percentage']}% ({analytics['group_bmb_grade']}) &nbsp;&nbsp;|&nbsp;&nbsp; "
        f"<b>Sana:</b> {datetime.now().strftime('%d.%m.%Y')}"
    )
    story.append(Paragraph(meta_text, meta_style))
    story.append(Spacer(1, 14))

    # BMB darajalar taqsimot paneli
    bmb = analytics["bmb_distribution"]
    bmb_summary = (
        f"<b>BMB Sertifikat Darajalari Taqsimoti:</b> &nbsp;&nbsp;"
        f"<b>A+ (Maksimal):</b> {bmb['A+']} nafar &nbsp;&nbsp;|&nbsp;&nbsp; "
        f"<b>A:</b> {bmb['A']} &nbsp;&nbsp;|&nbsp;&nbsp; "
        f"<b>B+:</b> {bmb['B+']} &nbsp;&nbsp;|&nbsp;&nbsp; "
        f"<b>B:</b> {bmb['B']} &nbsp;&nbsp;|&nbsp;&nbsp; "
        f"<b>C+:</b> {bmb['C+']} &nbsp;&nbsp;|&nbsp;&nbsp; "
        f"<b>C:</b> {bmb['C']} &nbsp;&nbsp;|&nbsp;&nbsp; "
        f"<b>Yetarli emas:</b> {bmb['Fail']}"
    )
    story.append(Paragraph(bmb_summary, meta_style))
    story.append(Spacer(1, 10))

    # Jadval
    table_data = [
        [
            Paragraph("№", th_style),
            Paragraph("O'quvchi F.I.Sh.", th_style),
            Paragraph("Email", th_style),
            Paragraph("Testlar", th_style),
            Paragraph("O'rtacha %", th_style),
            Paragraph("BMB Daraja", th_style),
            Paragraph("Esselar", th_style),
            Paragraph("Esse bali", th_style),
            Paragraph("Oxirgi faollik", th_style),
        ]
    ]

    for item in analytics["students"]:
        table_data.append(
            [
                Paragraph(str(item["rank"]), td_bold),
                Paragraph(item["full_name"], td_style),
                Paragraph(item["email"], td_style),
                Paragraph(str(item["tests_taken"]), td_bold),
                Paragraph(f"{item['avg_test_pct']}%", td_bold),
                Paragraph(item["bmb_grade"], td_bold),
                Paragraph(str(item["essays_written"]), td_bold),
                Paragraph(f"{item['avg_essay_score']}", td_bold),
                Paragraph(item["last_active"], td_style),
            ]
        )

    col_widths = [
        1.2 * cm,
        5.5 * cm,
        4.5 * cm,
        1.8 * cm,
        2.4 * cm,
        2.6 * cm,
        1.8 * cm,
        2.2 * cm,
        4.0 * cm,
    ]
    t = Table(table_data, colWidths=col_widths, repeatRows=1)
    t.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0F172A")),
                ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
                (
                    "ROWBACKGROUNDS",
                    (0, 1),
                    (-1, -1),
                    [colors.white, colors.HexColor("#F8FAFC")],
                ),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ]
        )
    )

    story.append(t)
    doc.build(story)
    return buf.getvalue()
